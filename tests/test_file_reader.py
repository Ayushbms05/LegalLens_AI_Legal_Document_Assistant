"""Tests for the file_reader utility."""

import io
from pathlib import Path
import sys
import pytest
from docx import Document

# Ensure the legallens root directory is on sys.path
LEGALLENS_DIR = Path(__file__).resolve().parent.parent
if str(LEGALLENS_DIR) not in sys.path:
    sys.path.insert(0, str(LEGALLENS_DIR))

from utils.file_reader import extract_text


class MockUploadedFile:
    """Helper class to mock a Streamlit UploadedFile object."""

    def __init__(self, name: str, data: bytes):
        self.name = name
        self._buffer = io.BytesIO(data)

    def read(self, *args):
        return self._buffer.read(*args)

    def seek(self, offset, whence=io.SEEK_SET):
        return self._buffer.seek(offset, whence)

    def tell(self):
        return self._buffer.tell()



def test_extract_text_txt():
    """Test extracting text from a plain text file."""
    sample_content = "This is a test legal agreement.\nClause 1: Payment is due immediately."
    mock_file = MockUploadedFile("contract.txt", sample_content.encode("utf-8"))

    extracted = extract_text(mock_file)
    assert "This is a test legal agreement." in extracted
    assert "Clause 1: Payment is due immediately." in extracted


def test_extract_text_docx():
    """Test extracting text from a .docx file."""
    doc = Document()
    doc.add_heading("Non-Disclosure Agreement", level=1)
    doc.add_paragraph("The receiving party agrees to maintain strict confidentiality.")

    # Save to an in-memory buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    mock_file = MockUploadedFile("nda.docx", buffer.getvalue())
    extracted = extract_text(mock_file)

    assert "Non-Disclosure Agreement" in extracted
    assert "The receiving party agrees to maintain strict confidentiality." in extracted


def test_extract_text_unsupported_format():
    """Test that an unsupported file format raises a clear ValueError."""
    mock_file = MockUploadedFile("spreadsheet.csv", b"col1,col2\nval1,val2")

    with pytest.raises(ValueError, match="Unsupported file format"):
        extract_text(mock_file)


def test_extract_text_empty_file():
    """Test that an empty file raises a ValueError."""
    mock_file = MockUploadedFile("empty.txt", b"   \n  ")

    with pytest.raises(ValueError, match="empty or contains no readable text"):
        extract_text(mock_file)
