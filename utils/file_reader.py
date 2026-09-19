"""File reader utility for LegalLens.

Extracts plain text from PDF, DOCX, and TXT files for legal document analysis.
"""

import io
from pathlib import Path
from typing import BinaryIO, Union
import warnings

from docx import Document
import pypdf

from utils.security import sanitize_input



MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15 MB maximum
MAX_PDF_PAGES = 100  # 100 pages maximum for DoS prevention


def extract_text(uploaded_file: Union[BinaryIO, io.BytesIO, str, Path]) -> str:
    """Extract clean plain text from a supported file (.pdf, .docx, .txt).

    Args:
        uploaded_file: A Streamlit UploadedFile, file-like object with a .name attribute,
                       or a file path (str or Path).

    Returns:
        Clean, extracted plain text as a string.

    Raises:
        ValueError: If the file format is unsupported, the file is empty,
                    the file exceeds size/page limits, or no readable text could be extracted.
    """
    # Step 1: Determine the filename and extension
    filename = getattr(uploaded_file, "name", None)
    if filename is None and isinstance(uploaded_file, (str, Path)):
        filename = str(uploaded_file)

    if not filename:
        raise ValueError("Could not determine the file name or file type.")

    # Validate file size if size attribute is present
    file_size = getattr(uploaded_file, "size", None)
    if file_size is not None and file_size > MAX_FILE_SIZE_BYTES:
        raise ValueError(
            f"File '{Path(filename).name}' exceeds the maximum allowed size of 15 MB."
        )

    extension = Path(filename).suffix.lower()

    # Step 2: Validate supported file extensions
    supported_extensions = {".pdf", ".docx", ".txt"}
    if extension not in supported_extensions:
        raise ValueError(
            f"Unsupported file format '{extension}'. "
            f"Please upload a .pdf, .docx, or .txt file."
        )

    # Step 3: Handle file reading based on extension
    extracted_text = ""

    # Ensure we have a file-like stream or open the path if given a string/Path
    file_stream = uploaded_file
    if isinstance(uploaded_file, (str, Path)):
        file_stream = open(uploaded_file, "rb")

    try:
        # Reset stream position to beginning if possible
        if hasattr(file_stream, "seek"):
            file_stream.seek(0)

        # -------------------------------------------------------------
        # TXT Extraction
        # -------------------------------------------------------------
        if extension == ".txt":
            raw_bytes = file_stream.read()
            if isinstance(raw_bytes, str):
                extracted_text = raw_bytes
            else:
                # Try UTF-8 decoding, with fallback to latin-1 for compatibility
                try:
                    extracted_text = raw_bytes.decode("utf-8")
                except UnicodeDecodeError:
                    extracted_text = raw_bytes.decode("latin-1", errors="replace")

        # -------------------------------------------------------------
        # DOCX Extraction (using python-docx)
        # -------------------------------------------------------------
        elif extension == ".docx":
            if isinstance(uploaded_file, (str, Path)):
                doc = Document(uploaded_file)
            else:
                # python-docx requires a seekable, tellable binary stream
                docx_bytes = file_stream.read()
                doc = Document(io.BytesIO(docx_bytes))

            text_parts = []


            # Extract text from all standard paragraphs
            for paragraph in doc.paragraphs:
                cleaned_para = paragraph.text.strip()
                if cleaned_para:
                    text_parts.append(cleaned_para)

            # Also extract text inside tables (common in contracts)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_parts.append(" | ".join(row_text))

            extracted_text = "\n\n".join(text_parts)

        # -------------------------------------------------------------
        # PDF Extraction (using pypdf)
        # -------------------------------------------------------------
        elif extension == ".pdf":
            reader = pypdf.PdfReader(file_stream)
            total_pages = len(reader.pages)

            if total_pages == 0:
                raise ValueError("The uploaded PDF has 0 pages.")

            if total_pages > MAX_PDF_PAGES:
                raise ValueError(
                    f"The uploaded PDF has {total_pages} pages, exceeding the {MAX_PDF_PAGES}-page limit."
                )

            page_texts = []
            for page_num, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    page_texts.append(page_text.strip())

            extracted_text = "\n\n".join(page_texts)

            # Warn if the document appears to be scanned (little to no extractable text)
            # A typical page of digital legal text has at least 150-200 characters.
            total_chars = len(extracted_text.strip())
            avg_chars_per_page = total_chars / total_pages if total_pages > 0 else 0
            if avg_chars_per_page < 50:
                warnings.warn(
                    "This PDF appears to be a scanned document or image. "
                    "Very little text could be extracted. Consider using OCR or a digital copy.",
                    UserWarning,
                    stacklevel=2,
                )

    finally:
        # Close file if we opened it from a string/Path
        if isinstance(uploaded_file, (str, Path)) and hasattr(file_stream, "close"):
            file_stream.close()

    # Step 4: Validate that the extracted text is not empty and sanitize it
    cleaned_output = sanitize_input(extracted_text.strip())

    if not cleaned_output:
        raise ValueError(
            f"The file '{Path(filename).name}' is empty or contains no readable text."
        )

    return cleaned_output
